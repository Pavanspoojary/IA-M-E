"""Unit tests for Excel workbook generation and openpyxl formatting."""

import pytest
import io
import openpyxl

from src.excel_exporter import ExcelExporter
from src.review_manager import StudentAssessmentRecord
from src.marks_validator import ValidatedMark, VerificationStatus
from src.template_manager import TemplateManager


@pytest.fixture
def exporter():
    return ExcelExporter()


@pytest.fixture
def sample_records():
    m1 = ValidatedMark("q1", "Q1", 5.0, 0.5, "4", 4.0, 4.0, 0.95, VerificationStatus.VALIDATED, "OK")
    m2 = ValidatedMark("q2", "Q2", 5.0, 0.5, "5", 5.0, 5.0, 0.90, VerificationStatus.CORRECTED_MANUALLY, "Manual")
    m3 = ValidatedMark("q3", "Q3", 5.0, 0.5, "", None, None, 1.0, VerificationStatus.BLANK, "Blank cell")

    r1 = StudentAssessmentRecord(
        source_filename="student1.png",
        page_number=2,
        student_id="1MS21CS042",
        student_name="Aarav Sharma",
        subject="Computer Networks",
        assessment="IA-1",
        marks={"q1": m1, "q2": m2, "q3": m3},
        total_marks=9.0,
        is_fully_verified=True,
    )
    return [r1]


def test_excel_workbook_structure(exporter, sample_records):
    """Test required worksheets exist with correct headers."""
    wb = exporter.generate_workbook(sample_records)
    sheet_names = wb.sheetnames

    assert "Marks Summary" in sheet_names
    assert "Detailed Extraction" in sheet_names
    assert "Processing Log" in sheet_names

    ws_summary = wb["Marks Summary"]
    assert ws_summary.cell(row=1, column=1).value == "Sl. No."
    assert ws_summary.cell(row=1, column=2).value == "Student ID"
    assert ws_summary.cell(row=2, column=2).value == "1MS21CS042"

    # Check that total formula or value is present
    total_val = ws_summary.cell(row=2, column=ws_summary.max_column - 1).value
    assert "=SUM" in str(total_val) or total_val == 9.0


def test_excel_byte_export(exporter, sample_records):
    """Test in-memory byte export can be re-opened by openpyxl."""
    raw_bytes = exporter.export_to_bytes(sample_records)
    assert len(raw_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
    assert len(wb.sheetnames) == 3


def test_detailed_extraction_content(exporter, sample_records):
    """Test detailed extraction sheet includes question rows and status highlights."""
    wb = exporter.generate_workbook(sample_records)
    ws = wb["Detailed Extraction"]

    # Header is row 1, 3 questions -> 3 rows of data
    assert ws.max_row == 4
    # Check recognition method column (col 8)
    assert "ONNX" in str(ws.cell(row=2, column=8).value)
    assert "Manual" in str(ws.cell(row=3, column=8).value)
