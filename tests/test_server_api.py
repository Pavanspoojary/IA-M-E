"""Automated tests for FastAPI web server endpoints.

Verifies:
1. Static and root index serving.
2. System status and CPU execution provider info.
3. Template retrieval and updates.
4. Batch processing and sample data loading.
5. Interactive mark updates with verification reconciliation.
6. Student metadata updates.
7. Excel export streaming.
"""

from pathlib import Path
import io
import pytest
from fastapi.testclient import TestClient

from server import app

BASE_DIR = Path(__file__).resolve().parent.parent


@pytest.fixture
def client():
    return TestClient(app)


def test_serve_index(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "IA Bluebook Marks Extraction System" in response.text
    assert "text/html" in response.headers["content-type"]


def test_api_status(client):
    response = client.get("/api/status")
    assert response.status_code == 200
    data = response.json()
    assert data["execution_provider"] == "CPUExecutionProvider"
    assert data["model_available"] is True
    assert data["total_configured_questions"] == 8
    assert data["total_max_marks"] == 60.0


def test_api_template(client):
    response = client.get("/api/template")
    assert response.status_code == 200
    data = response.json()
    assert "questions" in data
    assert len(data["questions"]) == 8


def test_load_samples_and_review_flow(client):
    # 1. Trigger sample loading
    resp = client.post(
        "/api/load-samples",
        data={
            "assessment": "IA Test - 1",
            "subject": "Operating Systems (21CS42)",
            "auto_align": "true",
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "records" in data
    assert len(data["records"]) == 4
    assert "summary" in data

    summary = data["summary"]
    assert summary["total_records"] == 4
    assert summary["total_fields"] == 32

    # 2. Pick the first record and inspect
    first_record = data["records"][0]
    record_id = first_record["record_id"]

    # 3. Update a mark manually (e.g. q1)
    update_resp = client.post(
        "/api/update-mark",
        json={
            "record_id": record_id,
            "question_id": "q1",
            "value": "4.5",
        },
    )
    assert update_resp.status_code == 200
    updated_data = update_resp.json()
    rec = updated_data["record"]
    assert rec["marks"]["q1"]["verified_val"] == 4.5
    assert rec["marks"]["q1"]["status"] == "Corrected manually"

    # 4. Update student metadata
    student_resp = client.post(
        "/api/update-student",
        json={
            "record_id": record_id,
            "student_id": "1MS21CS099",
            "student_name": "Test Student",
        },
    )
    assert student_resp.status_code == 200
    stud_data = student_resp.json()
    assert stud_data["record"]["student_id"] == "1MS21CS099"
    assert stud_data["record"]["student_name"] == "Test Student"

    # 5. Export Excel
    export_resp = client.get("/api/export")
    assert export_resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in export_resp.headers["content-type"]
    assert len(export_resp.content) > 1000


def test_extract_eval_sheet_endpoint(client):
    """Test the single-sheet 3-column IA extraction API."""
    resp = client.post(
        "/api/extract-sheet",
        data={
            "target_ia": "IA 1",
            "use_demo": "true",
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "column_crop_base64" in data
    assert "results_by_ia" in data
    assert "IA 1" in data["results_by_ia"]
    assert len(data["results_by_ia"]["IA 1"]) == 10

    totals = data["summary_totals"]
    assert "IA 1" in totals
    assert totals["IA 1"]["calculated_sum"] == 22.0
    assert totals["IA 1"]["detected_total"] == 22.0
    assert totals["IA 1"]["matches"] is True


def test_export_eval_sheet_endpoints(client):
    """Test Excel and CSV download for evaluation sheets."""
    # 1. Test CSV export with matrix format and UTF-8 BOM
    csv_resp = client.post(
        "/api/export-sheet",
        json={"format": "csv", "data": {}},
    )
    assert csv_resp.status_code == 200
    assert "text/csv" in csv_resp.headers["content-type"]
    assert csv_resp.content.startswith(b"\xef\xbb\xbf")
    assert "Sheet,Assessment,1a,1b,2a,2b,3a,3b,3b-alt,4a,4b,Total,Status" in csv_resp.text

    # 2. Test XLSX export
    xlsx_resp = client.post(
        "/api/export-sheet",
        json={"format": "xlsx", "data": {}},
    )
    assert xlsx_resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in xlsx_resp.headers["content-type"]
    assert len(xlsx_resp.content) > 1000


def test_extract_multipage_pdf_endpoint(client):
    """Test extracting marks from a multi-page PDF containing multiple mark sheets."""
    pdf_path = BASE_DIR / "sample_data" / "multipage_eval_sheets.pdf"
    assert pdf_path.exists(), "Multipage test PDF should exist"

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    resp = client.post(
        "/api/extract-sheet",
        files={"file": ("multipage_eval_sheets.pdf", pdf_bytes, "application/pdf")},
        data={"target_ia": "IA 1", "use_demo": "false"},
    )
    assert resp.status_code == 200
    data = resp.json()

    assert data["total_sheets"] == 2
    assert "sheets" in data
    assert len(data["sheets"]) == 2

    # Check Sheet 1
    sheet1 = data["sheets"][0]
    assert sheet1["sheet_index"] == 1
    assert sheet1["page_number"] == 1
    assert "IA 1" in sheet1["results_by_ia"]
    assert len(sheet1["results_by_ia"]["IA 1"]) == 10
    assert sheet1["column_crop_base64"].startswith("data:image/png;base64,")

    # Check Sheet 2
    sheet2 = data["sheets"][1]
    assert sheet2["sheet_index"] == 2
    assert sheet2["page_number"] == 2
    assert "IA 1" in sheet2["results_by_ia"]
    assert len(sheet2["results_by_ia"]["IA 1"]) == 10
    assert sheet2["column_crop_base64"].startswith("data:image/png;base64,")

    # Test Multi-Sheet CSV Export
    csv_resp = client.post(
        "/api/export-sheet",
        json={"format": "csv", "data": data},
    )
    assert csv_resp.status_code == 200
    assert "text/csv" in csv_resp.headers["content-type"]
    assert csv_resp.content.startswith(b"\xef\xbb\xbf")
    assert "Sheet,Assessment,1a,1b,2a,2b,3a,3b,3b-alt,4a,4b,Total,Status" in csv_resp.text
    assert "Sheet 1 (Page 1),IA 1" in csv_resp.text
    assert "Sheet 2 (Page 2),IA 1" in csv_resp.text

    # Test Multi-Sheet XLSX Export
    xlsx_resp = client.post(
        "/api/export-sheet",
        json={"format": "xlsx", "data": data},
    )
    assert xlsx_resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in xlsx_resp.headers["content-type"]
    assert len(xlsx_resp.content) > 1000

    # Verify openpyxl can load the multi-sheet workbook with individual tabs
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert "IA Marks Summary" in wb.sheetnames
    assert "Sheet 1" in wb.sheetnames
    assert "Sheet 2" in wb.sheetnames


def test_extract_user_multipage_photo_pdf_all_ia(client):
    """Test extracting all 6 photographed marksheet pages from real student booklet PDF with target_ia='ALL'."""
    pdf_path = BASE_DIR / "sample_data" / "user_multipage_sample.pdf"
    if not pdf_path.exists():
        pytest.skip("user_multipage_sample.pdf not found")

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    resp = client.post(
        "/api/extract-sheet",
        files={"file": ("user_multipage_sample.pdf", pdf_bytes, "application/pdf")},
        data={"target_ia": "ALL", "use_demo": "false"},
    )
    assert resp.status_code == 200
    data = resp.json()

    assert data["total_sheets"] == 6
    assert len(data["sheets"]) == 6

    # Verify all 6 sheets extracted all 3 IA columns
    for idx, sheet in enumerate(data["sheets"]):
        assert sheet["sheet_index"] == idx + 1
        assert "IA 1" in sheet["results_by_ia"]
        assert "IA 2" in sheet["results_by_ia"]
        assert "IA 3" in sheet["results_by_ia"]
        assert len(sheet["results_by_ia"]["IA 1"]) == 10
        assert len(sheet["results_by_ia"]["IA 2"]) == 10
        assert len(sheet["results_by_ia"]["IA 3"]) == 10

    # Test full 6-sheet Excel export
    xlsx_resp = client.post(
        "/api/export-sheet",
        json={"format": "xlsx", "data": data},
    )
    assert xlsx_resp.status_code == 200
    assert len(xlsx_resp.content) > 5000

    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert "IA Marks Summary" in wb.sheetnames
    for s_idx in range(1, 7):
        assert f"Sheet {s_idx}" in wb.sheetnames


def test_extract_distribution_of_marks_template(client):
    """Test extracting from Indian VTU/College 'DISTRIBUTION OF MARKS' bluebook template."""
    sample_path = BASE_DIR / "sample_data" / "distribution_marks_sample_2.jpg"
    if not sample_path.exists():
        pytest.skip("distribution_marks_sample_2.jpg not found")

    with open(sample_path, "rb") as f:
        img_bytes = f.read()

    resp = client.post(
        "/api/extract-sheet",
        files={"file": ("distribution_marks_sample_2.jpg", img_bytes, "image/jpeg")},
        data={"target_ia": "ALL", "use_demo": "false"},
    )
    assert resp.status_code == 200
    data = resp.json()

    assert data.get("template_type") == "DISTRIBUTION_OF_MARKS"
    assert "IA 1" in data["results_by_ia"]
    assert "IA 2" in data["results_by_ia"]
    assert "IA 3" in data["results_by_ia"]

    # Each IA has 13 rows (1a-4c + total)
    ia1_marks = data["results_by_ia"]["IA 1"]
    assert len(ia1_marks) == 13
    assert [m["id"] for m in ia1_marks] == [
        "1a", "1b", "1c", "2a", "2b", "2c", "3a", "3b", "3c", "4a", "4b", "4c", "total"
    ]

    # Verify detected digits
    ia1_map = {m["id"]: m for m in ia1_marks}
    assert ia1_map["1a"]["numeric_val"] == 6.0
    assert ia1_map["1b"]["numeric_val"] == 4.0
    assert ia1_map["total"]["numeric_val"] == 24.0

    # Verify CSV export with 12 questions
    csv_resp = client.post(
        "/api/export-sheet",
        json={"format": "csv", "data": data},
    )
    assert csv_resp.status_code == 200
    assert "Sheet,Assessment,1.a,1.b,1.c,2.a,2.b,2.c,3.a,3.b,3.c,4.a,4.b,4.c,Total,Status" in csv_resp.text

    # Verify XLSX export
    xlsx_resp = client.post(
        "/api/export-sheet",
        json={"format": "xlsx", "data": data},
    )
    assert xlsx_resp.status_code == 200
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert "IA Marks Summary" in wb.sheetnames



