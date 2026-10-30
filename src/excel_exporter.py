"""Excel workbook export module using openpyxl.

Produces formatted multi-worksheet examination marks workbooks with
custom styling, formulas, auto-filters, frozen panes, and status color coding.
"""

from typing import List, Optional
from pathlib import Path
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.review_manager import StudentAssessmentRecord
from src.marks_validator import VerificationStatus
from src.template_manager import BluebookTemplate


class ExcelExporter:
    """Exports structured assessment records into openpyxl workbooks."""

    # Header styling
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Borders
    THIN_BORDER = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Status fills
    STATUS_FILLS = {
        VerificationStatus.VALIDATED: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        VerificationStatus.VERIFIED_AUTOMATICALLY: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        VerificationStatus.VERIFIED_MANUALLY: PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
        VerificationStatus.CORRECTED_MANUALLY: PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
        VerificationStatus.NEEDS_REVIEW: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        VerificationStatus.INVALID: PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        VerificationStatus.UNREADABLE: PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        VerificationStatus.BLANK: PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
    }

    def generate_workbook(
        self,
        records: List[StudentAssessmentRecord],
        template: Optional[BluebookTemplate] = None,
    ) -> openpyxl.Workbook:
        """Generate a complete 3-sheet Excel workbook.

        Args:
            records: List of StudentAssessmentRecord items.
            template: Optional BluebookTemplate containing ordered question labels.

        Returns:
            openpyxl.Workbook instance.
        """
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Determine question columns list
        question_keys: List[str] = []
        question_labels: List[str] = []

        if template and template.questions:
            for q in template.questions:
                question_keys.append(q.id)
                question_labels.append(q.label)
        elif records and records[0].marks:
            for q_id, mark in records[0].marks.items():
                question_keys.append(q_id)
                question_labels.append(mark.question_label)

        # Create Worksheets
        self._build_marks_summary_sheet(wb, records, question_keys, question_labels)
        self._build_detailed_extraction_sheet(wb, records)
        self._build_processing_log_sheet(wb, records)

        return wb

    def export_to_bytes(
        self,
        records: List[StudentAssessmentRecord],
        template: Optional[BluebookTemplate] = None,
    ) -> bytes:
        """Export workbook directly to memory bytes for web downloads."""
        wb = self.generate_workbook(records, template)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    def export_to_file(
        self,
        records: List[StudentAssessmentRecord],
        output_path: str | Path,
        template: Optional[BluebookTemplate] = None,
    ) -> None:
        """Save formatted workbook to disk."""
        wb = self.generate_workbook(records, template)
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))

    def _build_marks_summary_sheet(
        self,
        wb: openpyxl.Workbook,
        records: List[StudentAssessmentRecord],
        q_keys: List[str],
        q_labels: List[str],
    ) -> None:
        """Worksheet 1: Marks Summary."""
        ws = wb.create_sheet(title="Marks Summary")

        headers = [
            "Sl. No.",
            "Student ID",
            "Student Name",
            "Subject",
            "Assessment",
        ] + q_labels + [
            "Total Marks",
            "Verification Status",
        ]

        ws.append(headers)

        for sl_no, rec in enumerate(records, start=1):
            row_data = [
                sl_no,
                rec.student_id,
                rec.student_name,
                rec.subject,
                rec.assessment,
            ]

            # Append question marks
            first_q_col = 6
            last_q_col = first_q_col + len(q_keys) - 1

            for q_k in q_keys:
                mark = rec.marks.get(q_k)
                if mark and mark.verified_val is not None:
                    row_data.append(float(mark.verified_val))
                elif mark and mark.status == VerificationStatus.BLANK:
                    row_data.append("")
                else:
                    # Unverified or invalid: leave empty or flag
                    row_data.append(None)

            # Total marks: write Excel SUM formula if question columns exist
            row_idx = sl_no + 1
            if q_keys:
                start_letter = get_column_letter(first_q_col)
                end_letter = get_column_letter(last_q_col)
                total_formula = f"=SUM({start_letter}{row_idx}:{end_letter}{row_idx})"
                row_data.append(total_formula)
            else:
                row_data.append(rec.total_marks if rec.total_marks is not None else 0.0)

            status_str = "Fully Verified" if rec.is_fully_verified else "Requires Review"
            row_data.append(status_str)

            ws.append(row_data)

        self._style_worksheet(ws, status_col_idx=len(headers))

    def _build_detailed_extraction_sheet(
        self,
        wb: openpyxl.Workbook,
        records: List[StudentAssessmentRecord],
    ) -> None:
        """Worksheet 2: Detailed Extraction."""
        ws = wb.create_sheet(title="Detailed Extraction")

        headers = [
            "Student ID",
            "Source Filename",
            "PDF Page Number",
            "Question Label",
            "Raw Extracted Value",
            "Final Verified Value",
            "Maximum Marks",
            "Recognition Method",
            "Verification Status",
        ]
        ws.append(headers)

        for rec in records:
            for q_id, mark in rec.marks.items():
                method_desc = (
                    "Manual Teacher Correction"
                    if mark.status in (VerificationStatus.CORRECTED_MANUALLY, VerificationStatus.VERIFIED_MANUALLY)
                    else f"ONNX CPU Model ({rec.model_version})"
                )

                v_val = float(mark.verified_val) if mark.verified_val is not None else ""
                row_data = [
                    rec.student_id,
                    rec.source_filename,
                    rec.page_number,
                    mark.question_label,
                    mark.extracted_str or "",
                    v_val,
                    float(mark.max_marks),
                    method_desc,
                    mark.status,
                ]
                ws.append(row_data)

        self._style_worksheet(ws, status_col_idx=9)

    def _build_processing_log_sheet(
        self,
        wb: openpyxl.Workbook,
        records: List[StudentAssessmentRecord],
    ) -> None:
        """Worksheet 3: Processing Log."""
        ws = wb.create_sheet(title="Processing Log")

        headers = [
            "Source Filename",
            "Processing Timestamp",
            "Page-Processing Result",
            "Model Version",
            "Fields Requiring Review",
            "Error or Warning Message",
        ]
        ws.append(headers)

        for rec in records:
            res_str = "Success" if rec.is_fully_verified else "Requires Review"
            review_count = rec.needs_review_count + rec.invalid_count

            row_data = [
                rec.source_filename,
                rec.processing_timestamp,
                res_str,
                rec.model_version,
                review_count,
                rec.log_message,
            ]
            ws.append(row_data)

        self._style_worksheet(ws, status_col_idx=3)

    def _style_worksheet(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        status_col_idx: Optional[int] = None,
    ) -> None:
        """Apply professional styling, filters, borders, and column autosizing."""
        # Header Row Styling
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = self.HEADER_ALIGN

        # Freeze top header row
        ws.freeze_panes = "A2"

        # Apply Auto-Filter
        last_col_letter = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"

        # Style data rows
        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = self.THIN_BORDER
                cell.font = Font(name="Calibri", size=10)

                # Center short codes and status
                if col == 1 or col == status_col_idx:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                # Status color highlight
                if col == status_col_idx and cell.value in self.STATUS_FILLS:
                    cell.fill = self.STATUS_FILLS[cell.value]

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
